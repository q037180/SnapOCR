"""文本导出：按扩展名将识别/翻译结果保存为 txt / docx / pdf。

pdf 走 PySide6 自带的 QPrinter（中文渲染无障碍）；
docx 走 python-docx（延迟导入，未安装时给出友好提示）。
"""
from __future__ import annotations

import os


def export_text(path: str, text: str) -> None:
    """按扩展名导出文本。支持 .txt / .docx / .pdf。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".txt":
        _save_txt(path, text)
    elif ext == ".docx":
        _save_docx(path, text)
    elif ext == ".pdf":
        _save_pdf(path, text)
    else:
        raise ValueError(f"不支持的导出格式：{ext}")


def _save_txt(path: str, text: str) -> None:
    # utf-8-sig：带 BOM，Windows 记事本/Word 打开不乱码
    with open(path, "w", encoding="utf-8-sig") as f:
        f.write(text)


def _save_docx(path: str, text: str) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "导出 docx 需要 python-docx：请先执行 pip install python-docx"
        ) from exc
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(path)


def _save_pdf(path: str, text: str) -> None:
    from PySide6.QtGui import QTextDocument
    from PySide6.QtPrintSupport import QPrinter

    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(path)
    doc = QTextDocument()
    html = "".join(
        (line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") or "<br>")
        + "<br>"
        for line in text.splitlines()
    )
    doc.setHtml(html)
    doc.print_(printer)
